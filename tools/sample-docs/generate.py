"""Generate deterministic, entirely fictional supplier documents. No external assets."""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import os
import random
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFilter, ImageFont

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs/pitch/contoh-dokumen-supplier"
SEED = 690926
STAMP = datetime(2026, 9, 30, 17, 0)
MATERIALS = ["Split 1-2", "Screening 5-10", "Split 2-3", "Abu Batu 0-5"]
SHEETS = ["SPLITE 1-2", "SCREENING 5-10", "SPLITE 2-3", "ABU BATU 0-5"]
LOCS = {"CKD": "Cikande", "SERANG": "Serang"}
SUPPLIER = "PT. BATU ALAM SEJAHTERA"
CUSTOMER = "PT. Beton Cipta Mandiri"
ADMIN = "RATMA"
SENDER = "admin.timbangan.bas@gmail.com"
RECIPIENT = "rekap@nsa-demo.example"
DRIVERS = ["BAYU", "JAKA", "DANU", "TAMA", "RUDI", "YUDI", "DIMAS", "GALIH", "NANDA", "ARIF", "WIRA", "RAKA"]
PLATES = [f"B{9131 + i * 37}Q{chr(65 + i)}Z" for i in range(len(DRIVERS))]
HEAD4 = ["No", "No", "No", "No", "No", "SUPIR", "Tanggal", "Isi(kg)", "Kosong(Kg)", "Qty(kg)", "Material", "Keterangan", "No"]
HEAD5 = [None, "GJIB", "DO", "Surat Jalan", "Polisi", None, None, None, None, None, None, None, "PO"]


def font(size, kind="mono"):
    candidates = {
        "mono": ["DejaVuSansMono.ttf", "consola.ttf", "LiberationMono-Regular.ttf"],
        "sans": ["segoeui.ttf", "DejaVuSans.ttf", "LiberationSans-Regular.ttf"],
        "bold": ["segoeuib.ttf", "DejaVuSans-Bold.ttf", "LiberationSans-Bold.ttf"],
        "hand": ["segoepr.ttf", "DejaVuSans-Oblique.ttf", "LiberationSans-Italic.ttf"],
    }[kind]
    roots = [Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts", Path("/usr/share/fonts/truetype/dejavu"), Path("/usr/share/fonts/truetype/liberation2"), Path("/System/Library/Fonts")]
    for folder in roots:
        for name in candidates:
            path = folder / name
            if path.exists():
                return ImageFont.truetype(str(path), size)
    return ImageFont.load_default(size=size)


def kg(value):
    return f"{value:,}".replace(",", ".")


def fields(r):
    return {
        "supplier": SUPPLIER, "bap": f"{r['bap']:06d}", "tanggal": r["date"],
        "waktu_masuk": f"{r['date']} {r['in_time']}", "nomor_do": str(r["do"]),
        "waktu_keluar": f"{r['date']} {r['out_time']}", "customer": "NSA",
        "project": f"{CUSTOMER}- {LOCS[r['location']]}", "no_truck": r["plate"],
        "driver": r["driver"], "product": r["material"] + " AA",
        "qty_kg": r["net_kg"], "berat_kosong_kg": r["tare_kg"],
        "berat_isi_kg": r["gross_kg"], "sales_order_no": r["po_full"],
        "keterangan": "NSA " + ("CKD" if r["location"] == "CKD" else "SRG"),
        "no_sj": r["sj"], "tujuan": f"{CUSTOMER} - {LOCS[r['location']]}",
        "barang": r["material"], "admin": ADMIN, "signature_driver_name": r["driver"],
        "signature_security_name": "GENTA", "signature_recipient_name": "LARAS",
    }


def build_data():
    rng = random.Random(SEED)
    rows, pos = [], []
    next_history_po, next_september_po = 1, 69
    for month in (7, 8, 9):
        for li, loc in enumerate(LOCS):
            sizes = [11, 11, 11, 11, 18] if month != 9 else [21, 11, 11, 11, 11, 11, 11, 11, 12]
            pool = []
            for i, size in enumerate(sizes):
                pnum = next_september_po if month == 9 else next_history_po
                if month == 9:
                    next_september_po += 1
                else:
                    next_history_po += 1
                roman = {7: "VII", 8: "VIII", 9: "IX"}[month]
                p = {"number": pnum, "full": f"PO-{pnum:03d}/BS/NSA/S/{roman}/2026", "material": MATERIALS[(i + li) % 4], "location": loc, "planned_trucks": size, "planned_tons": 700 if size >= 18 else 350}
                pos.append(p)
                pool.append([p, size])
            cursor = 0
            days = 30 if month == 9 else 31
            for day in range(1, days + 1):
                count = (3 if day % 3 == 1 else 4) if month == 9 else 2
                if month == 9 and day in (21, 22):
                    count = 3 if day == 21 else 4
                for trip in range(count):
                    while pool[cursor % len(pool)][1] == 0:
                        cursor += 1
                    p, _ = pool[cursor % len(pool)]
                    pool[cursor % len(pool)][1] -= 1
                    cursor += 1
                    driver_idx = (day + trip + li * 3 + month) % len(DRIVERS)
                    if month == 9 and day == 22 and loc == "CKD" and trip < 2:
                        driver_idx = 2
                    hour = 7 + trip * 3
                    minute = rng.randrange(8, 32)
                    net = rng.randint(2900, 4100) * 10
                    tare = rng.randint(1140, 1280) * 10
                    rows.append({"date": date(2026, month, day).isoformat(), "location": loc, "daily_no": trip + 1, "in_time": f"{hour:02d}:{minute:02d}", "out_time": f"{hour:02d}:{minute+21:02d}", "plate": PLATES[driver_idx], "driver": DRIVERS[driver_idx], "material": p["material"], "po": p["number"], "po_full": p["full"], "net_kg": net, "tare_kg": tare, "gross_kg": net + tare, "remark": "BETON " + LOCS[loc].upper()})
            assert all(n == 0 for _, n in pool)
    rows.sort(key=lambda r: (r["date"], r["in_time"], r["location"]))
    for i, row in enumerate(rows):
        row.update(id=f"TRIP-{i+1:04d}", do=24100+i, bap=24261+i, sj=460+i)
    def pick(loc, day, trip):
        return next(r for r in rows if r["location"] == loc and r["date"] == f"2026-09-{day:02d}" and r["daily_no"] == trip)
    cases = {
        "excel_vs_do_3_tons": pick("CKD", 22, 1)["id"],
        "unclear_do": pick("SERANG", 22, 2)["id"],
        "excel_without_photo": pick("CKD", 23, 1)["id"],
        "photo_without_excel": pick("SERANG", 23, 4)["id"],
        "billed_weight_changed": pick("CKD", 5, 1)["id"],
        "wrong_location_remark": pick("CKD", 23, 2)["id"],
        "two_trips_one_day": [pick("CKD", 22, 1)["id"], pick("CKD", 22, 2)["id"]],
    }
    old = next(r for r in rows if r["id"] == cases["billed_weight_changed"])
    old["net_kg"] = 30000
    old["gross_kg"] = 30000 + old["tare_kg"]
    for p in pos:
        rr = [r for r in rows if r["po"] == p["number"]]
        p["actual_trucks"] = len(rr)
        p["actual_tons"] = round(sum(r["net_kg"] for r in rr)/1000, 2)
    return rows, pos, cases


def excel_weights(r, version, cases):
    change = 3000 if r["id"] == cases["excel_vs_do_3_tons"] else 0
    if version == 2 and r["id"] == cases["billed_weight_changed"]:
        change = 5000
    return (r["gross_kg"] + change)/1000, r["tare_kg"]/1000, (r["net_kg"] + change)/1000


def canonical_xlsx(path):
    # ZIP timestamps and document properties must also be deterministic.
    with zipfile.ZipFile(path) as src:
        entries = [(i.filename, src.read(i.filename)) for i in src.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as dst:
        for name, payload in entries:
            info = zipfile.ZipInfo(name, (2026, 9, 30, 17, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            if name == "docProps/core.xml":
                import re
                payload = re.sub(rb'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)', rb'\g<1>2026-09-30T17:00:00Z\2', payload)
            dst.writestr(info, payload)


def write_excel(rows, cases, location, version):
    name = f"REKAP NSA-{location}" + (" (2)" if version == 2 else "") + ".xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Admin Timbangan BAS (fiktif)"
    wb.properties.created = wb.properties.modified = STAMP
    refs = {}
    border = Border(*( [Side(style="thin", color="777777")] * 4))
    cutoff = "2026-09-23" if version == 1 else "2026-09-30"
    for material, sheet in zip(MATERIALS, SHEETS):
        ws = wb.create_sheet(sheet)
        for col in range(1, 14):
            ws.cell(4, col, HEAD4[col-1])
            if HEAD5[col-1]:
                ws.cell(5, col, HEAD5[col-1])
            else:
                ws.merge_cells(start_row=4, start_column=col, end_row=5, end_column=col)
            for rr in (4, 5):
                c = ws.cell(rr, col)
                c.font = Font(name="Arial", size=10, bold=True)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                c.fill = PatternFill("solid", fgColor="D9E5F2")
                c.border = border
        selected = [r for r in rows if r["location"] == location and r["material"] == material and r["date"] <= cutoff and r["id"] != cases["photo_without_excel"]]
        daily = defaultdict(list)
        for row in selected:
            daily[row["date"]].append(row)
        totals = []
        for day, trips in sorted(daily.items()):
            dayqty = []
            for day_no, r in enumerate(trips, 1):
                gross, tare, qty = excel_weights(r, version, cases)
                remark = "BETON SERANG" if r["id"] == cases["wrong_location_remark"] else r["remark"]
                # Numbering is per material/day, as each material has its own sheet.
                ws.append([day_no, r["bap"], r["do"], r["sj"], r["plate"], r["driver"], datetime.fromisoformat(day), gross, tare, qty, material, remark, r["po"]])
                refs[r["id"]] = {"file": name, "sheet": sheet, "row": ws.max_row, "gross_tons": gross, "tare_tons": tare, "qty_tons": qty, "remark": remark, "po_sequence": r["po"]}
                ws.cell(ws.max_row, 7).number_format = "dd/mm/yyyy"
                for cc in (8, 9, 10):
                    ws.cell(ws.max_row, cc).number_format = "0.00"
                dayqty.append(round(qty*100))
            subtotal = sum(dayqty)/100
            ws.append([None]*9 + [subtotal])
            ws.cell(ws.max_row, 10).number_format = "0.00"
            ws.cell(ws.max_row, 10).font = Font(bold=True)
            totals.extend(dayqty)
        total = sum(totals)/100
        # Keep footer totals outside identifier columns A-F; the ERP parser
        # otherwise interprets C/D as an extra DO/SJ. Label spans G:H.
        ws.append([None]*6 + ["SUBTOTAL", None, None, total])
        ws.append([None]*6 + [material, None, "=", total, "Ton"])
        ws.merge_cells(start_row=ws.max_row, start_column=7, end_row=ws.max_row, end_column=8)
        ws.cell(ws.max_row, 10).number_format = "0.00"
        ws.cell(ws.max_row-1, 10).number_format = "0.00"
        widths = [6, 11, 11, 15, 16, 13, 14, 12, 13, 12, 20, 24, 9]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "G6"
        ws.row_dimensions[4].height = ws.row_dimensions[5].height = 22
        ws.sheet_view.showGridLines = True
        ws.print_options.horizontalCentered = True
        ws.print_title_rows = "4:5"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.center.text = "DOKUMEN FIKTIF - DEMO NSA"
    wb.save(OUT / name)
    canonical_xlsx(OUT / name)
    return refs


def signature(draw, box, name, rng, ink):
    x, y, w, h = box
    points = []
    phase = rng.uniform(0, 2)
    amplitude = rng.uniform(10, 20)
    for i in range(180):
        u = i/179
        px = x+12 + u*(w-50) + 7*math.sin(u*math.pi*22+phase)
        py = y+h*.42 - amplitude*math.sin(u*math.pi*12+phase)*(1-u*.35) - 9*u
        points.append((px, py))
    draw.line(points, fill=ink, width=2)
    draw.line([(x+18, y+h*.63), (x+w-22, y+h*.50)], fill=ink, width=2)
    draw.text((x+15, y+h-34), name.title(), font=font(23, "hand"), fill=ink)


def do_form(r, rng):
    w, h = 1500, 900
    paper = Image.new("RGB", (w, h), (241, 241, 223))
    pd = ImageDraw.Draw(paper)
    for x in (22, w-23):
        for y in range(24, h, 48):
            pd.ellipse((x-8, y-8, x+8, y+8), fill=(64, 64, 59))
    pd.line((45, 0, 45, h), fill=(191, 192, 177), width=1)
    pd.line((w-46, 0, w-46, h), fill=(191, 192, 177), width=1)
    mask = Image.new("L", (w, h))
    d = ImageDraw.Draw(mask)
    ff = font(25)
    def t(x, y, s, size=None):
        d.text((x, y), s, font=font(size) if size else ff, fill=255)
    t(76, 45, SUPPLIER, 30)
    d.line((76, 84, 555, 84), fill=255, width=2)
    t(1120, 45, "BAP: " + f"{r['bap']:06d}", 27)
    dt = date.fromisoformat(r["date"]).strftime("%d/%m/%Y")
    left = [
        ("Tanggal", dt), ("Waktu Masuk", dt + " " + r["in_time"]),
        ("Customer", "NSA"), ("Project", f"{CUSTOMER}- {LOCS[r['location']]}"),
        ("No Truck", r["plate"]), ("Product", r["material"] + " AA"),
        ("Berat Kosong", kg(r["tare_kg"])), ("Sales Order No", r["po_full"]),
        ("Keterangan", "NSA " + ("CKD" if r["location"] == "CKD" else "SRG")),
    ]
    right = {0: ("Nomor DO", str(r["do"])), 1: ("Waktu Keluar", dt + " " + r["out_time"]), 4: ("Driver", r["driver"]), 5: ("Qty (Kg)", kg(r["net_kg"])), 6: ("Berat Isi (Kg)", kg(r["gross_kg"]))}
    for i, (label, value) in enumerate(left):
        y = 118 + i*53
        t(76, y, label)
        t(300, y, ": " + value)
        if i in right:
            label, value = right[i]
            t(810, y, label)
            t(1050, y, ": " + value)
    # Simulated dot-matrix raster: pale ink, regular pin grid, missing pin strikes.
    m = np.asarray(mask).astype(np.float32)/255
    yy, xx = np.indices((h, w))
    pins = (((xx % 3)-1)**2 + ((yy % 3)-1)**2 <= 1)
    breaks = rng.random()
    nr = np.random.default_rng(r["do"] + SEED)
    coverage = m * pins * (nr.random((h, w)) > .035 - breaks*.01)
    arr = np.asarray(paper).astype(np.float32)
    arr -= coverage[..., None] * rng.uniform(94, 116)
    paper = Image.fromarray(np.uint8(np.clip(arr, 0, 255)))
    d = ImageDraw.Draw(paper)
    for i, label in enumerate(("Admin", "Driver", "Security", "Penerima Barang")):
        x = 76 + i*335
        d.rectangle((x, 650, x+320, 838), outline=(165, 165, 154), width=2)
        d.text((x+14, 664), label, font=font(24), fill=(139, 139, 131))
        if i == 0:
            d.text((x+14, 850), ADMIN, font=font(22), fill=(137, 137, 128))
        else:
            signature(d, (x, 700, 315, 128), ["", r["driver"], "GENTA", "LARAS"][i], rng, (90, 102, 118))
    d.text((1020, 861), "DOKUMEN FIKTIF / DEMO", font=font(16), fill=(159, 159, 148))
    return paper


def sj_form(r, rng, carbon):
    paper = Image.new("RGB", (1500, 900), (244, 202, 211) if carbon else (244, 239, 221))
    d = ImageDraw.Draw(paper)
    printed = (129, 73, 93) if carbon else (92, 98, 104)
    ink = (137, 84, 121) if carbon else (41, 68, 120)
    d.text((90, 46), "NSA", font=font(48, "bold"), fill=printed)
    d.text((90, 105), "SURAT JALAN", font=font(34, "bold"), fill=printed)
    d.text((1110, 65), "No. " + str(r["sj"]), font=font(37, "hand"), fill=ink)
    d.line((80, 161, 1420, 161), fill=printed, width=2)
    vals = [("Tanggal", date.fromisoformat(r["date"]).strftime("%d/%m/%Y")), ("No Polisi", r["plate"]), ("Sopir", r["driver"].title()), ("Tujuan", f"{CUSTOMER} - {LOCS[r['location']]}")]
    for i, (label, value) in enumerate(vals):
        y = 184 + i*61
        d.text((90, y), label, font=font(27, "sans"), fill=printed)
        d.text((280+rng.randint(-4, 4), y-5+rng.randint(-3, 3)), value, font=font(32, "hand"), fill=ink)
    d.rectangle((80, 447, 1420, 637), outline=printed, width=2)
    for x in (155, 915, 1150):
        d.line((x, 447, x, 637), fill=printed, width=2)
    d.line((80, 503, 1420, 503), fill=printed, width=2)
    for x, text in [(100, "No"), (185, "Barang"), (952, "Jumlah"), (1180, "Satuan")]:
        d.text((x, 462), text, font=font(27, "sans"), fill=printed)
    for x, text in [(106, "1"), (182, r["material"]), (945, f"{r['net_kg']/1000:.2f}".replace(".", ",")), (1200, "Ton")]:
        d.text((x, 536), text, font=font(34, "hand"), fill=ink)
    for i, label in enumerate(("Pengirim", "Sopir", "Security", "Penerima")):
        x = 80 + i*338
        d.rectangle((x, 673, x+326, 839), outline=printed, width=2)
        d.text((x+17, 682), label, font=font(25, "sans"), fill=printed)
        signature(d, (x, 720, 320, 118), [ADMIN, r["driver"], "GENTA", "LARAS"][i], rng, ink)
    d.text((88, 856), "Lembar " + ("2 - salinan karbon" if carbon else "1 - penerima"), font=font(18, "sans"), fill=printed)
    d.text((1110, 856), "FIKTIF / DEMO", font=font(18), fill=printed)
    return paper


def perspective_coeffs(dest, src):
    a, b = [], []
    for (x, y), (u, v) in zip(dest, src):
        a.extend([[x, y, 1, 0, 0, 0, -u*x, -u*y], [0, 0, 0, x, y, 1, -v*x, -v*y]])
        b.extend([u, v])
    return np.linalg.solve(np.asarray(a), np.asarray(b)).tolist()


def phone_photo(r, kind, index, poor):
    rng = random.Random(SEED + r["do"]*13 + (1 if kind == "SJ" else 0))
    nr = np.random.default_rng(SEED + r["do"]*23 + index)
    carbon = kind == "SJ" and index % 3 != 0
    paper = sj_form(r, rng, carbon) if kind == "SJ" else do_form(r, rng)
    pw, ph = paper.size
    # Creased paper: local bend shifts rows and shades the fold.
    a = np.asarray(paper).copy()
    for y in range(ph):
        shift = round(3*math.sin(y/95) + 2*math.sin(y/29))
        a[y] = np.roll(a[y], shift, axis=0)
    crease = rng.randrange(300, 1100)
    cx = np.arange(pw)
    fold = 1 - .08*np.exp(-((cx-crease)/8)**2) + .04*np.exp(-((cx-crease-10)/7)**2)
    a = np.clip(a.astype(float)*fold[None, :, None], 0, 255).astype(np.uint8)
    paper = Image.fromarray(a).convert("RGBA")
    # The form is photographed sideways, alternating clockwise/counterclockwise.
    clockwise = index % 2 == 0
    paper = paper.transpose(Image.Transpose.ROTATE_270 if clockwise else Image.Transpose.ROTATE_90)
    sw, sh = paper.size
    x = rng.randint(39, 72)
    top = rng.randint(74, 115)
    dest = [(x+rng.randint(-12, 10), top), (850+rng.randint(-18, 12), top+rng.randint(-28, 28)), (840+rng.randint(-25, 10), 1515+rng.randint(-25, 20)), (x+rng.randint(-12, 22), 1490+rng.randint(-15, 35))]
    bg = Image.new("RGBA", (900, 1600), (36, 43, 41, 255))
    bd = ImageDraw.Draw(bg)
    bd.line((0, 1250, 900, 1390), fill=(53, 62, 57), width=5)
    # Sheets below the main form: uneven edges, offset layers and visible shadow.
    for dx, dy in [(18, 24), (-15, 12), (9, 8)]:
        corners = [(px+dx, py+dy) for px, py in dest]
        bd.polygon(corners, fill=(193+dy, 191+dy, 178+dy, 255))
    shadow = Image.new("RGBA", bg.size)
    ImageDraw.Draw(shadow).polygon([(px+15, py+20) for px, py in dest], fill=(0, 0, 0, 110))
    bg = Image.alpha_composite(bg, shadow.filter(ImageFilter.GaussianBlur(15)))
    coeffs = perspective_coeffs(dest, [(0, 0), (sw, 0), (sw, sh), (0, sh)])
    warped = paper.transform(bg.size, Image.Transform.PERSPECTIVE, coeffs, Image.Resampling.BICUBIC)
    bg = Image.alpha_composite(bg, warped)
    a = np.asarray(bg.convert("RGB")).astype(np.float32)
    yy, xx = np.indices((1600, 900), dtype=np.float32)
    # Like the owner's phone photos: a washed-out sheet (dull white, pale ink),
    # light falling unevenly from one side, the phone's shadow, a dark glass desk
    # that reflects the ceiling lamps, a green-grey cast, visible sensor noise.
    a = a*.84 + 22
    side = 1 if rng.random() < .5 else -1
    vignette = 1 - .22*((xx-450)/620)**2 - .16*((yy-800)/1000)**2
    illumination = rng.uniform(.9, 1.04) + side*.09*(xx-450)/450 + .05*(yy-800)/800
    phone_shadow = .22*np.exp(-(((xx-rng.randrange(150, 750))/300)**2 + ((yy-rng.randrange(1250, 1500))/300)**2))
    a *= (vignette*illumination-phone_shadow)[..., None]
    # Lamps reflected in the glass, mostly beside the sheet, one spilling onto it.
    for _ in range(3):
        gx = rng.choice([rng.uniform(5, 70), rng.uniform(830, 895)])
        gy = rng.uniform(150, 1450)
        amp = rng.uniform(70, 140)
        a += (amp*np.exp(-(((xx-gx)/rng.uniform(45, 80))**2 + ((yy-gy)/rng.uniform(60, 120))**2)))[..., None]
    spill_x = rng.uniform(120, 780)
    a += (rng.uniform(18, 34)*np.exp(-(((xx-spill_x)/220)**2 + ((yy-rng.uniform(300, 1300))/260)**2)))[..., None]
    a *= np.array([.97, 1.0, .95], dtype=np.float32)
    a += nr.normal(0, 4.2, a.shape[:2]).astype(np.float32)[..., None]
    a += nr.normal(0, 1.6, a.shape).astype(np.float32)
    image = Image.fromarray(np.clip(a, 0, 255).astype(np.uint8))
    obscured = None
    if poor:
        # Right-hand weight cells in the unrotated form; map into sideways photo.
        if clockwise:
            box = (340, 1060, 530, 1460)
        else:
            box = (345, 170, 555, 565)
        image.paste(image.crop(box).filter(ImageFilter.GaussianBlur(6)), box)
        overlay = Image.new("RGBA", image.size)
        ImageDraw.Draw(overlay).rounded_rectangle(box, radius=70, fill=(8, 13, 18, 130))
        image = Image.alpha_composite(image.convert("RGBA"), overlay.filter(ImageFilter.GaussianBlur(28))).convert("RGB")
        obscured = {"bbox_xyxy": list(box), "fields": ["qty_kg", "berat_isi_kg"], "effect": "bayangan tangan/HP dan blur lokal"}
    # Small exposure/optical softness, sensor noise and JPEG compression.
    image = image.filter(ImageFilter.GaussianBlur(rng.uniform(.65, 1.0)))
    path = f"foto/{kind}_{r['do']:06d}_{r['location']}.jpg"
    image.save(OUT / path, quality=70, subsampling=2, optimize=True)
    return {"file": path, "kind": kind, "trip_id": r["id"], "fields": fields(r), "photo": {"size": [900, 1600], "rotation_degrees": 90 if clockwise else -90, "perspective_corners": dest, "carbon_copy": carbon, "poor_quality": poor, "obscured": obscured}, "intentional_differences": []}


def email_screenshot(batch, files, count):
    im = Image.new("RGB", (1500, 1020), "#ffffff")
    d = ImageDraw.Draw(im)
    def txt(x, y, text, size=23, color="#303842", kind="sans"):
        d.text((x, y), text, font=font(size, kind), fill=color)
    d.rectangle((0, 0, 1500, 58), fill="#0f6cbd")
    txt(24, 13, "Outlook", 25, "white", "bold")
    txt(1040, 18, "SIMULASI EMAIL / DATA FIKTIF", 17, "white")
    d.rectangle((0, 59, 1500, 119), fill="#f5f6f8")
    txt(28, 77, "File     Beranda     Tampilan     Bantuan", 21)
    txt(605, 77, "Balas       Balas semua       Teruskan", 21)
    d.rectangle((0, 120, 222, 1020), fill="#f1f3f5")
    txt(24, 146, "Favorit", 23, kind="bold")
    d.rectangle((10, 191, 214, 236), fill="#d9eaf8")
    txt(25, 200, "Kotak Masuk   2", 22)
    txt(25, 258, "Item Terkirim", 22)
    txt(25, 315, "Draf", 22)
    txt(252, 144, batch["subject"], 33, kind="bold")
    d.ellipse((254, 220, 308, 274), fill="#0f6cbd")
    txt(262, 230, "RT", 22, "white")
    txt(325, 214, "Ratma - Admin Timbangan BAS", 24, kind="bold")
    txt(325, 251, SENDER, 21, "#59636e")
    txt(252, 298, "Kepada: " + RECIPIENT, 22)
    txt(252, 335, batch["date_display"] + " WIB", 20, "#69727d")
    txt(252, 394, f"{count} lampiran  |  2 rekap Excel dan {count-2} foto DO/SJ", 23, kind="bold")
    previews = files[:2] + [f for f in files[2:] if "/DO_" in f][:2] + [f for f in files[2:] if "/SJ_" in f][:2]
    for i, name in enumerate(previews):
        x, y = 254+(i % 3)*392, 441+(i//3)*94
        d.rounded_rectangle((x, y, x+373, y+78), radius=5, fill="#f5f7f9", outline="#d2d8df")
        txt(x+15, y+12, "XLSX" if name.endswith("xlsx") else "JPG", 18, "#107c41" if name.endswith("xlsx") else "#0f6cbd", "bold")
        txt(x+15, y+42, Path(name).name, 18)
    txt(253, 644, f"+ {count-6} lampiran lainnya", 20, "#0f6cbd")
    for i, line in enumerate(batch["body"].splitlines()):
        txt(252, 704 + i*34, line, 23)
    im.save(OUT / batch["screenshot"], optimize=True)


def write_emails(photo_entries):
    batches = []
    for v, start, end in [(1, "2026-09-22", "2026-09-23"), (2, "2026-09-29", "2026-09-30")]:
        tag = "22-230926" if v == 1 else "29-300926"
        suffix = " (2)" if v == 2 else ""
        excel = [f"REKAP NSA-{loc}{suffix}.xlsx" for loc in LOCS]
        photos = [p["file"] for p in photo_entries if start <= p["fields"]["tanggal"] <= end]
        rng = random.Random(SEED+v)
        rng.shuffle(photos)
        files = excel + photos
        body = "Selamat sore,\n\nTerlampir rekap serta DO dan surat jalan NSA untuk Serang dan Cikande.\nMohon dicek ya. Terima kasih.\n\nRatma\nAdmin Timbangan BAS"
        msg = EmailMessage(policy=policy.SMTP)
        msg["From"] = f"Ratma - Admin Timbangan BAS <{SENDER}>"
        msg["To"] = RECIPIENT
        msg["Subject"] = f"Rekap & SJ NSA-SRG,CKD {tag}"
        day = date.fromisoformat(end)
        msg["Date"] = f"{day:%a, %d %b %Y} 17:05:00 +0700"
        msg["Message-ID"] = f"<nsa-supplier-demo-{v}-202609@bas-demo.example>"
        msg["X-Demo-Fictional"] = "true"
        msg.set_content(body, charset="utf-8")
        for f in files:
            is_excel = f.endswith("xlsx")
            msg.add_attachment((OUT/f).read_bytes(), maintype="application" if is_excel else "image", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet" if is_excel else "jpeg", filename=Path(f).name)
        msg.set_boundary(f"==NSA-FIKTIF-{SEED}-KIRIMAN-{v}==")
        eml = f"email/Kiriman-{tag}.eml"
        (OUT/eml).write_bytes(msg.as_bytes())
        batch = {"version": v, "file": eml, "screenshot": f"email/Outlook-{tag}.png", "subject": str(msg["Subject"]), "sender": SENDER, "recipient": RECIPIENT, "date": end+"T17:05:00+07:00", "date_display": day.strftime("%d/%m/%Y")+" 17:05", "body": body, "photo_range": [start, end], "attachments": files, "image_count": len(photos)}
        email_screenshot(batch, files, len(files))
        batches.append(batch)
    return batches


def make_manifest(rows, pos, cases, refs, photos, emails):
    trips = {r["id"]: r for r in rows}
    c = []
    descriptions = {
        "excel_vs_do_3_tons": ("Qty rekap 3 ton lebih besar daripada DO; Isi ikut naik agar Qty = Isi - Kosong.", "Tandai selisih berat foto vs rekap sebesar +3.00 ton."),
        "unclear_do": ("Angka Qty dan Isi foto DO tertutup bayangan serta blur.", "Minta pemeriksaan manual; jangan menganggap OCR sebagai kebenaran."),
        "excel_without_photo": ("Truk tercatat di Excel, tetapi foto DO dan SJ sengaja tidak disertakan.", "Tandai DO belum dilampirkan pada kiriman 22-23 September."),
        "photo_without_excel": ("Foto DO/SJ ada, baris truk dihilangkan dari kedua versi Excel.", "Tandai foto DO tanpa baris rekap."),
        "billed_weight_changed": ("Truk 5 September sudah ditagih; rekap berikutnya diam-diam mengubah Qty 30.00 menjadi 35.00 ton.", "Tandai perubahan pada transaksi yang sudah ditagih, jangan terapkan otomatis."),
        "wrong_location_remark": ("Baris berada di file Cikande tetapi Keterangan BETON SERANG.", "Tandai ketidaksesuaian lokasi."),
        "two_trips_one_day": ("Plat dan sopir sama mengambil dua kali pada 22 September untuk material dan PO berbeda.", "Pertahankan dua transaksi; jangan gabungkan berdasarkan tanggal + plat saja."),
    }
    for key, ids in cases.items():
        ids = ids if isinstance(ids, list) else [ids]
        case = {"case": key, "description": descriptions[key][0], "expected_behavior": descriptions[key][1], "trip_ids": ids, "do_numbers": [trips[t]["do"] for t in ids], "excel_refs": {t: {f"v{v}": refs[v].get(t) for v in (1, 2)} for t in ids}, "photos": [p["file"] for p in photos if p["trip_id"] in ids]}
        if key == "billed_weight_changed":
            case["application_fixture"] = {"trip_id": ids[0], "status": "Ditagih", "invoice_number": "INV-DEMO-NSA-202609-017", "invoice_date": "2026-09-18", "invoiced_qty_tons": 30.0, "revised_qty_tons": 35.0, "setup_required": "Sebelum impor v2, set transaksi ini Ditagih di aplikasi memakai nilai v1; generator tidak mengubah database ERP."}
        c.append(case)
        for p in photos:
            if p["trip_id"] in ids and not (key == "unclear_do" and p["kind"] != "DO"):
                p["intentional_differences"].append(key)
    for p in photos:
        if p["photo"]["poor_quality"] and "unclear_do" not in p["intentional_differences"]:
            p["intentional_differences"].append("additional_unclear_do")
        p["excel_refs"] = {f"v{v}": refs[v].get(p["trip_id"]) for v in (1, 2)}
    return {"schema_version": 1, "fictional": True, "seed": SEED, "generated_at": "2026-09-30T17:00:00+07:00", "supplier": SUPPLIER, "buyer": "NSA", "customer": CUSTOMER, "period": ["2026-07-01", "2026-09-30"], "bap_minus_do": 161, "snapshot_cutoffs": {"v1": "2026-09-23", "v2": "2026-09-30"}, "weight_units": {"excel": "ton, 2 desimal walau judul kg", "do": "kg, titik ribuan", "manifest_trip_weights": "integer kg"}, "notes": ["Semua data fiktif. Tidak ada referensi asli yang digunakan.", "Foto hanya untuk jendela kiriman 22-23 dan 29-30 September. Riwayat Excel tidak memerlukan foto di kiriman ini.", "Tata letak SJ perkiraan; foto dibuat dengan Pillow/numpy tanpa aset internet.", "Status Ditagih adalah fixture pengujian dalam manifest, bukan hasil mutasi database ERP.", "Nomor urut kolom No dimulai 1 per tanggal dalam masing-masing sheet material.", "PO 69 dan 78 menargetkan sekitar 700 ton dengan 21 truk; PO September lain 11-12 truk sekitar 350-420 ton."], "purchase_orders": pos, "trips": rows, "photos": photos, "test_cases": c, "emails": emails, "counts": {"actual_trips": len(rows), "excel_v1": len(refs[1]), "excel_v2": len(refs[2]), "photos": len(photos), "do_photos": sum(p["kind"] == "DO" for p in photos), "sj_photos": sum(p["kind"] == "SJ" for p in photos)}}


def validate(manifest):
    errors = []
    def check(ok, msg):
        if not ok:
            errors.append(msg)
    records = {r["id"]: r for r in manifest["trips"]}
    cases = {c["case"]: c for c in manifest["test_cases"]}
    case_ids = {k: v["trip_ids"][0] for k, v in cases.items()}
    check(len(cases) == 7, "Harus ada tujuh kasus uji")
    seen = {}
    for version in (1, 2):
        for loc in LOCS:
            suffix = " (2)" if version == 2 else ""
            name = f"REKAP NSA-{loc}{suffix}.xlsx"
            wb = load_workbook(OUT/name, data_only=True)
            check(wb.sheetnames == SHEETS, name+": urutan/nama sheet")
            extracted = {}
            for mat, ws in zip(MATERIALS, wb.worksheets):
                check(all(ws.cell(r, c).value is None for r in (1, 2, 3) for c in range(1, 14)), name+": baris 1-3 kosong")
                check([ws.cell(4, c).value for c in range(1, 14)] == HEAD4, name+": header baris 4")
                check([ws.cell(5, c).value for c in range(1, 14)] == HEAD5, name+": header baris 5")
                check({str(m) for m in ws.merged_cells.ranges if m.min_row == 4} == {f"{get_column_letter(c)}4:{get_column_letter(c)}5" for c in (1, 6, 7, 8, 9, 10, 11, 12)}, name+": merge vertikal")
                dayqty, grand, lastday, n = 0, 0, None, 0
                for values in ws.iter_rows(min_row=6, values_only=True):
                    if isinstance(values[0], int):
                        check(isinstance(values[6], datetime), name+": tanggal Excel asli")
                        check(values[10] == mat, name+": material sheet")
                        check(abs(values[7]-values[8]-values[9]) < .00001, name+": Qty != Isi - Kosong")
                        check(isinstance(values[12], int), name+": PO wajib angka")
                        check(" " not in values[4], name+": plat tanpa spasi")
                        thisday = values[6].date().isoformat()
                        n = n+1 if thisday == lastday else 1
                        lastday = thisday
                        check(values[0] == n, name+": urut harian")
                        dayqty += round(values[9]*100)
                        grand += round(values[9]*100)
                        check(values[1]-values[2] == 161, name+": BAP-DO")
                        extracted[values[2]] = values
                    elif values[6] == "SUBTOTAL":
                        check(dayqty == 0 and round(values[9]*100) == grand, name+": SUBTOTAL akhir")
                    elif values[6] == mat:
                        check(values[8] == "=" and values[10] == "Ton" and round(values[9]*100) == grand, name+": total material")
                    else:
                        check(all(v is None for i, v in enumerate(values) if i != 9), name+": subtotal hanya kolom Qty")
                        check(round(values[9]*100) == dayqty, name+": subtotal harian")
                        dayqty = 0
            check(all(v[6].date().isoformat() <= manifest["snapshot_cutoffs"][f"v{version}"] for v in extracted.values()), name+": cutoff")
            seen[(version, loc)] = extracted
            wb.close()
    v1 = {k: v for loc in LOCS for k, v in seen[(1, loc)].items()}
    v2 = {k: v for loc in LOCS for k, v in seen[(2, loc)].items()}
    check(len(v1) == manifest["counts"]["excel_v1"], "jumlah baris v1")
    check(len(v2) == manifest["counts"]["excel_v2"], "jumlah baris v2")
    changed = [do for do in v1 if v1[do] != v2[do]]
    old = records[case_ids["billed_weight_changed"]]
    check(changed == [old["do"]], "Hanya satu baris lama berubah")
    check(v1[old["do"]][9] == 30 and v2[old["do"]][9] == 35, "Perubahan 30 -> 35")
    diff = records[case_ids["excel_vs_do_3_tons"]]
    check(round(v1[diff["do"]][9]*1000)-diff["net_kg"] == 3000, "Selisih 3 ton")
    orphan = records[case_ids["photo_without_excel"]]
    check(orphan["do"] not in v1 and orphan["do"] not in v2, "Foto tanpa Excel")
    missing = records[case_ids["excel_without_photo"]]
    check(missing["do"] in v1 and not any(p["trip_id"] == missing["id"] for p in manifest["photos"]), "Excel tanpa foto")
    wrong = records[case_ids["wrong_location_remark"]]
    check(v1[wrong["do"]][11] == "BETON SERANG" and wrong["do"] in seen[(1, "CKD")], "Keterangan salah lokasi")
    same = [records[t] for t in cases["two_trips_one_day"]["trip_ids"]]
    check(same[0]["date"] == same[1]["date"] and same[0]["plate"] == same[1]["plate"] and same[0]["po"] != same[1]["po"] and same[0]["material"] != same[1]["material"], "Dua trip satu plat berbeda material/PO")
    daily = Counter((r["date"], r["location"]) for r in records.values() if r["date"] >= "2026-09-01")
    check(len(daily) == 60 and all(2 <= n <= 4 for n in daily.values()), "2-4 truk/hari/lokasi September")
    check(all(29000 <= r["net_kg"] <= 41000 and 11400 <= r["tare_kg"] <= 12800 for r in records.values()), "rentang berat")
    check([r["do"] for r in records.values()] == list(range(24100, 24100+len(records))), "DO berurutan")
    check([r["sj"] for r in records.values()] == list(range(460, 460+len(records))), "SJ berurutan lintas lokasi")
    for p in manifest["purchase_orders"]:
        rr = [r for r in records.values() if r["po"] == p["number"]]
        check(11 <= len(rr) <= 21 and len({(r["material"], r["location"]) for r in rr}) == 1, "Satu PO 11-21 truk, satu material/lokasi")
    for p in manifest["photos"]:
        path = OUT/p["file"]
        with Image.open(path) as im:
            im.verify()
        with Image.open(path) as im:
            check(im.size == (900, 1600) and im.format == "JPEG", p["file"]+": format foto")
        check(p["fields"] == fields(records[p["trip_id"]]), p["file"]+": fields manifest")
    poor = [p for p in manifest["photos"] if p["kind"] == "DO" and p["photo"]["poor_quality"]]
    check(len(poor) == 3 and any(p["trip_id"] == case_ids["unclear_do"] for p in poor), "Sekitar 1/10 DO kurang jelas")
    for batch in manifest["emails"]:
        msg = BytesParser(policy=policy.default).parsebytes((OUT/batch["file"]).read_bytes())
        parts = list(msg.iter_attachments())
        check(not msg.defects and str(msg["Subject"]) == batch["subject"], "EML parse/header")
        check([p.get_filename() for p in parts] == [Path(f).name for f in batch["attachments"]], "Daftar lampiran EML")
        for part, f in zip(parts, batch["attachments"]):
            check(part.get_payload(decode=True) == (OUT/f).read_bytes(), "Byte lampiran EML "+f)
        check(30 <= batch["image_count"] <= 40 and sum(p.get_content_type().endswith("sheet") for p in parts) == 2, "Campuran 30-40 foto + 2 Excel")
        with Image.open(OUT/batch["screenshot"]) as im:
            im.verify()
    total = sum(p.stat().st_size for p in OUT.rglob("*") if p.is_file())
    check(total < 30_000_000, f"Ukuran {total} melebihi 30 MB")
    if errors:
        raise RuntimeError("Validasi gagal:\n" + "\n".join(errors))
    return {"ok": True, "xlsx": 4, "photos": len(manifest["photos"]), "emails": 2, "screenshots": 2, "test_cases": 7, "bytes": total, "megabytes": round(total/1_000_000, 2), "scope": "openpyxl, Pillow, MIME attachments, subtotal, konsistensi data dan semua kasus; tidak menjalankan ERP atau OCR"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validate-only", action="store_true", help="Periksa keluaran yang sudah ada tanpa membuat ulang")
    args = parser.parse_args()
    if args.validate_only:
        manifest = json.loads((OUT/"manifest.json").read_text(encoding="utf-8"))
    else:
        for folder in (OUT, OUT/"foto", OUT/"email"):
            folder.mkdir(parents=True, exist_ok=True)
        rows, pos, cases = build_data()
        refs = {1: {}, 2: {}}
        for version in (1, 2):
            for loc in LOCS:
                refs[version].update(write_excel(rows, cases, loc, version))
        print("4 rekap Excel selesai", flush=True)
        selected = [r for r in rows if r["date"] in ("2026-09-22", "2026-09-23", "2026-09-29", "2026-09-30") and r["id"] != cases["excel_without_photo"]]
        # Keep precisely three unclear DOs out of 31, including the named test case.
        unclear = {cases["unclear_do"], next(r["id"] for r in selected if r["date"] == "2026-09-29" and r["location"] == "CKD"), next(r["id"] for r in selected if r["date"] == "2026-09-30" and r["location"] == "SERANG")}
        photos = []
        for i, row in enumerate(selected):
            for kind in ("DO", "SJ"):
                photos.append(phone_photo(row, kind, i, kind == "DO" and row["id"] in unclear))
            if (i+1) % 8 == 0:
                print(f"Foto {2*(i+1)}/{2*len(selected)} selesai", flush=True)
        emails = write_emails(photos)
        manifest = make_manifest(rows, pos, cases, refs, photos, emails)
        (OUT/"manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    report = validate(manifest)
    (OUT/"validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    # Ignore validation.json size in its own byte count so repeated runs match.
    report["bytes"] = sum(p.stat().st_size for p in OUT.rglob("*") if p.is_file() and p.name != "validation.json")
    report["megabytes"] = round(report["bytes"]/1_000_000, 2)
    (OUT/"validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    digest = hashlib.sha256()
    for p in sorted(OUT.rglob("*")):
        if p.is_file():
            digest.update(p.relative_to(OUT).as_posix().encode())
            digest.update(p.read_bytes())
    print(json.dumps(report, ensure_ascii=False), flush=True)
    print("SHA256 seluruh keluaran: " + digest.hexdigest(), flush=True)


if __name__ == "__main__":
    main()
